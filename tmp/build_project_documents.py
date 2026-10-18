from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docx"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
BLUE = "DDEBF7"
PALE = "F3F6FA"
MID = "5B6573"
GRID = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def configure_doc(doc, short_title):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.7)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title.element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 18, 14, 7),
        ("Heading 2", 14, 11, 5),
        ("Heading 3", 11.5, 8, 3),
    ):
        st = styles[name]
        st.font.name = "Aptos Display"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for section in doc.sections:
        section.different_first_page_header_footer = True
        hp = section.header.paragraphs[0]
        hp.text = short_title
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.runs[0].font.name = "Aptos"
        hp.runs[0].font.size = Pt(8.5)
        hp.runs[0].font.color.rgb = RGBColor(91, 101, 115)
        add_page_number(section.footer.paragraphs[0])
        section.first_page_header.paragraphs[0].text = ""
        section.first_page_footer.paragraphs[0].text = ""


def add_cover(doc, title, subtitle, date_line, audience_line):
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p.paragraph_format.space_before = Pt(72)
    ppr = p._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)

    sp = doc.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = sp.add_run(subtitle)
    r.font.name = "Aptos Display"
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(50, 50, 50)
    sp.paragraph_format.space_after = Pt(28)

    meta = doc.add_table(rows=3, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    values = [
        ("Document purpose", audience_line),
        ("Prepared", "12 October 2026"),
        ("Presentation date", date_line),
    ]
    for i, (label, value) in enumerate(values):
        set_cell_width(meta.cell(i, 0), 1.55)
        set_cell_width(meta.cell(i, 1), 5.0)
        meta.cell(i, 0).text = label
        meta.cell(i, 1).text = value
        for c in meta.rows[i].cells:
            set_cell_border(c, WHITE, "0")
            set_cell_margins(c, 80, 0, 80, 100)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        meta.cell(i, 0).paragraphs[0].runs[0].bold = True
        meta.cell(i, 0).paragraphs[0].runs[0].font.color.rgb = RGBColor(91, 101, 115)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    r = p.add_run("Weather driven decisions explained with transparent assumptions and reproducible analysis")
    r.font.name = "Aptos"
    r.font.size = Pt(11)
    r.font.italic = True
    r.font.color.rgb = RGBColor(91, 101, 115)
    doc.add_page_break()


def add_table(doc, headers, rows, widths=None, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    set_repeat_table_header(table.rows[0])
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(font_size)
        if widths:
            set_cell_width(cell, widths[j])
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, value in enumerate(row):
            cells[j].text = str(value)
            set_cell_border(cells[j])
            set_cell_margins(cells[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2 == 1:
                set_cell_shading(cells[j], PALE)
            for p in cells[j].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.size = Pt(font_size)
            if widths:
                set_cell_width(cells[j], widths[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Inches(0.24 + 0.2 * level)
        p.paragraph_format.first_line_indent = Inches(-0.16)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.add_run(item)


def add_process(doc, labels):
    table = doc.add_table(rows=1, cols=len(labels))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    width = 6.72 / len(labels)
    for i, label in enumerate(labels):
        c = table.cell(0, i)
        c.text = f"{i + 1}\n{label}"
        set_cell_width(c, width)
        set_cell_shading(c, BLUE if i % 2 == 0 else PALE)
        set_cell_border(c, GRID)
        set_cell_margins(c, 140, 80, 140, 80)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in c.paragraphs[0].runs:
            r.font.size = Pt(9)
            r.font.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph()
    r = p.add_run(label)
    r.bold = True
    p.add_run(text)
    return p


def build_report():
    doc = Document()
    configure_doc(doc, "Weather Driven Business Intelligence Platform")
    add_cover(
        doc,
        "Weather Driven Business Intelligence Platform",
        "Technical Project Report",
        "13 October 2026",
        "A clear and evidence based explanation of the project architecture methods results and limitations",
    )

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(
        "This project demonstrates how weather data can be combined with business operations data to estimate and communicate weather related commercial effects. It joins two years of synthetic hourly weather and business telemetry across four United States cities, controls for recurring calendar patterns, and presents the results through an interactive Streamlit dashboard. Live weather forecasts enrich the demonstration, while uploaded merchant data can be used to estimate business specific relationships."
    )
    doc.add_paragraph(
        "The principal conclusion is that the platform is a strong end to end analytical prototype. Its value lies in data alignment, transparent feature engineering, adjusted regression, scenario exploration, and accessible visualization. It should not be presented as a production forecasting service or as proof of causal business effects because the principal historical dataset is synthetic and the forward revenue projections are rule based."
    )

    add_table(
        doc,
        ["Project dimension", "Current implementation"],
        [
            ("Historical data", "Synthetic weather and cross sector operational telemetry for 2023 and 2024"),
            ("Geographic scope", "New York City Chicago Seattle and Phoenix"),
            ("Processed warehouse", "70045 synchronized city hour observations and 61 variables"),
            ("Analytical scope", "21 operational metrics across six economic sectors"),
            ("Live enrichment", "Open Meteo forecast data NASA POWER reference data and NASA Worldview links"),
            ("Interface", "Five tab Streamlit application with interactive Plotly charts"),
            ("Verification", "Seven automated tests passing in the reviewed environment"),
        ],
        widths=[1.65, 5.05],
    )

    doc.add_heading("1 Project purpose and business question", level=1)
    doc.add_paragraph(
        "Businesses experience weather differently. Rain can increase delivery demand while reducing restaurant footfall and bike usage. Heat can lift cold beverage sales while increasing electricity and HVAC demand. The platform turns these relationships into a portfolio view that asks three practical questions:"
    )
    add_numbered(
        doc,
        [
            "How do weather conditions relate to revenue and operational demand after recurring calendar patterns are considered",
            "How might one weather event propagate across connected sectors such as mobility dining delivery and freight",
            "What staffing inventory and operating actions should managers consider before forecast weather arrives",
        ],
    )
    doc.add_paragraph(
        "The system is designed as a demonstration of analytical workflow rather than a substitute for a company forecasting process. Its scenarios help users explore assumptions and prioritize questions for real data collection."
    )

    doc.add_heading("2 System architecture", level=1)
    add_process(doc, ["Generate raw feeds", "Align time and location", "Engineer features", "Estimate relationships", "Simulate impacts", "Present decisions"])
    doc.add_paragraph(
        "The architecture separates data creation, processing, modeling, business logic, external weather retrieval, and presentation. This modular structure makes the project easier to test and extend than a single notebook or monolithic dashboard."
    )
    add_table(
        doc,
        ["Component", "Responsibility", "Primary output"],
        [
            ("data generator", "Creates reproducible synthetic weather and operational records with seasonality and weather sensitivity", "Raw CSV feeds"),
            ("data pipeline", "Normalizes timestamps resamples business events imputes weather gaps and engineers analytical variables", "Merged Parquet warehouse"),
            ("econometrics engine", "Fits adjusted regression models and evaluates user defined weather scenarios", "Coefficients comparisons and scenario metrics"),
            ("revenue analytics", "Applies business catalog assumptions and creates portfolio ripple narratives", "Revenue series and operational guidance"),
            ("live weather", "Retrieves current and forecast conditions with resilient fallback behavior", "Seven day weather package"),
            ("Streamlit application", "Coordinates controls charts diagnostics and explanations", "Interactive dashboard"),
        ],
        widths=[1.25, 3.45, 2.0],
        font_size=8.8,
    )

    doc.add_heading("3 Data generation and integration", level=1)
    doc.add_heading("3.1 Source design", level=2)
    doc.add_paragraph(
        "The historical analysis uses two intentionally mismatched synthetic feeds. Weather observations arrive hourly in UTC with sensor gaps. Business events arrive every 15 minutes in local time and contain operational measures for delivery mobility retail hospitality energy and freight. The mismatch creates a realistic engineering exercise without implying that the records came from real merchants."
    )
    add_table(
        doc,
        ["Raw source", "Rows", "Selected fields", "Timing"],
        [
            ("Weather telemetry", "69106", "Temperature precipitation snow wind humidity and condition", "Hourly UTC"),
            ("Cross sector business stream", "279249", "Orders mobility retail dining energy freight and warehouse measures", "15 minute local time"),
        ],
        widths=[1.5, 0.72, 3.55, 0.93],
        font_size=9,
    )
    doc.add_heading("3.2 Alignment and feature engineering", level=2)
    add_bullets(
        doc,
        [
            "Local business timestamps are assigned the appropriate IANA timezone and converted to UTC.",
            "Fifteen minute business records are aggregated into hourly measures using sums or means according to the metric.",
            "Continuous weather gaps are filled using time based interpolation; missing precipitation and snow increments are set to zero.",
            "The weather and business tables are joined by city and UTC hour.",
            "Calendar controls include hour weekday month weekend and selected holiday indicators.",
            "Weather features include rain categories temperature categories one and two hour precipitation lags and temperature anomalies.",
            "Cross sector ratios measure courier utilization delivery versus dine in activity and rideshare versus bike usage.",
        ],
    )
    doc.add_paragraph(
        "The resulting warehouse contains 70045 observations with no missing fields and no duplicated city hour keys in the reviewed output. Temperature anomaly is calculated relative to the average for the same city month and hour in the available dataset."
    )

    doc.add_heading("4 Analytical method", level=1)
    doc.add_heading("4.1 Why raw correlation is insufficient", level=2)
    doc.add_paragraph(
        "Weather and commerce both follow recurring cycles. A rainy Friday evening can coincide with high delivery demand even when Friday itself explains much of the increase. Comparing raw correlation with an adjusted model shows why calendar patterns should be considered before attributing an operational change to weather."
    )
    doc.add_heading("4.2 Regression specification", level=2)
    doc.add_paragraph(
        "For each target metric the project fits an ordinary least squares model containing precipitation snow temperature anomaly wind and categorical controls for hour weekday month holiday status and city. HC1 heteroskedasticity robust standard errors are used. The implementation is best described as OLS with city and calendar fixed effects."
    )
    add_label_paragraph(
        doc,
        "Interpretation. ",
        "A precipitation coefficient estimates the average within model change in the target associated with one additional millimeter of precipitation while the included controls are held constant. Percentage effects are created by dividing the coefficient by the target mean.",
    )
    add_label_paragraph(
        doc,
        "Important limitation. ",
        "These estimates are adjusted associations. The model does not by itself prove causality and the synthetic generator already embeds many of the relationships the regression later recovers.",
    )
    add_table(
        doc,
        ["Reported quantity", "What it means", "What it does not mean"],
        [
            ("Coefficient", "Estimated conditional change in the target for a one unit change in a weather variable", "Guaranteed causal effect"),
            ("P value", "Compatibility of the estimate with a null hypothesis under the model assumptions", "Probability that the effect is true"),
            ("R squared", "Share of in sample variation explained by the fitted model", "Out of sample prediction accuracy"),
            ("Scenario delta", "Difference between two model predictions under specified inputs", "Observed future revenue"),
        ],
        widths=[1.2, 2.8, 2.7],
        font_size=9,
    )

    doc.add_heading("5 Business sectors and scenario logic", level=1)
    add_table(
        doc,
        ["Sector", "Illustrative metrics", "Expected weather pathway"],
        [
            ("Delivery and logistics", "Orders couriers surge delivery time cancellations", "Rain can increase demand while reducing available courier supply"),
            ("Urban mobility", "Bike trips transit rideshare congestion", "Travelers substitute away from exposed transport modes"),
            ("Retail and commerce", "Footfall rainwear winter apparel ecommerce", "Weather changes store visits product mix and online demand"),
            ("Hospitality", "Dine in guests patio use hot food cold beverages", "Conditions alter venue choice and product preference"),
            ("Energy", "Grid load and HVAC demand", "Temperature extremes increase heating or cooling requirements"),
            ("Freight", "Road delay and warehouse dispatch", "Poor conditions can reduce transport reliability and throughput"),
        ],
        widths=[1.45, 2.5, 2.75],
        font_size=8.8,
    )
    doc.add_paragraph(
        "The dashboard converts selected weather conditions into sector predictions and business specific revenue illustrations. Domino chains describe plausible propagation paths such as rain reducing bike use increasing rideshare demand worsening congestion and extending delivery times. These chains are useful scenario narratives but are not estimated structural causal networks."
    )

    doc.add_heading("6 Dashboard experience", level=1)
    add_table(
        doc,
        ["Dashboard view", "Purpose"],
        [
            ("Weekly Operational Impact", "Compares modeled actual revenue with a clear weather baseline for the selected portfolio"),
            ("Historical Trends and Seasonality", "Shows longer term revenue patterns and compares raw correlations with adjusted model estimates"),
            ("Business Interactions and Domino Chains", "Explains cross sector substitutions thresholds and downstream operational effects"),
            ("Live Scenario Simulator and Guidance", "Uses current or user selected weather conditions and displays a seven day weather adjusted projection"),
            ("Data Wrangling Diagnostics", "Exposes alignment logic data volumes engineered variables and synchronized records"),
        ],
        widths=[2.25, 4.45],
    )
    doc.add_paragraph(
        "Users may choose preset businesses create a custom category or upload CSV or Excel revenue history. Custom business parameters are inferred from keywords or supplied directly. Uploaded merchant revenue is aligned with the project weather history to estimate business specific sensitivities when enough overlapping dates exist."
    )

    doc.add_heading("7 Live data and forecasting boundaries", level=1)
    add_bullets(
        doc,
        [
            "Open Meteo supplies the current conditions and seven day meteorological forecast used by the application.",
            "NASA POWER supplies a recent historical reference record rather than a same moment confirmation of current conditions.",
            "NASA Worldview is presented as an external imagery browser link and is not analyzed automatically by the model.",
            "Square and Stripe connector classes can call external APIs when credentials are supplied, but the current application initializes them without credentials and therefore uses calibrated sandbox sequences.",
            "Seven day revenue projections apply configured elasticities and weekday weights to weather forecasts. They are deterministic scenario projections rather than trained machine learning or generative AI forecasts.",
        ],
    )

    doc.add_heading("8 Verification and current quality", level=1)
    doc.add_paragraph(
        "The project includes unit and integration checks for data generation pipeline alignment scenario output connector schemas CSV column detection and merchant elasticity estimation. All seven tests passed during this review. The processed warehouse was also checked for missing values and duplicate city hour keys."
    )
    add_table(
        doc,
        ["Verified item", "Result"],
        [
            ("Automated tests", "7 passed"),
            ("Processed observations", "70045"),
            ("Processed variables", "61"),
            ("Missing values", "0 in the processed warehouse"),
            ("Duplicate city hour keys", "0"),
            ("Historical period", "1 January 2023 through 31 December 2024 subject to city timezone boundaries"),
        ],
        widths=[2.15, 4.55],
    )

    doc.add_heading("9 Limitations and recommended next steps", level=1)
    add_table(
        doc,
        ["Priority", "Current limitation", "Recommended improvement"],
        [
            ("High", "Historical relationships are generated rather than observed", "Validate with permissioned merchant and operational data from multiple businesses"),
            ("High", "Adjusted regressions are described too strongly as causal", "Use careful association language and develop a credible identification strategy before causal claims"),
            ("High", "Default Square and Stripe records are simulated", "Add explicit credential configuration status labels pagination and visible connector errors"),
            ("High", "Merchant model fit is constrained to a preset range", "Report the actual R squared and add standard errors confidence intervals and holdout validation"),
            ("Medium", "Seven day revenue projections are deterministic", "Add backtesting uncertainty intervals and benchmark models"),
            ("Medium", "Weather sources represent different observation times", "Display source timestamps prominently and avoid treating delayed NASA data as a current consensus check"),
            ("Medium", "Dashboard supports four selectable cities", "Add geocoding and configurable city metadata if broader coverage is required"),
        ],
        widths=[0.65, 2.55, 3.5],
        font_size=8.5,
    )

    doc.add_page_break()
    doc.add_heading("10 Conclusion", level=1)
    doc.add_paragraph(
        "The Weather Driven Business Intelligence Platform is a coherent demonstration of multi source data engineering econometric adjustment scenario analysis and executive visualization. It succeeds as a portfolio project because the complete analytical path is visible from raw data through operational recommendations. Its credibility is strongest when the synthetic nature of the history and the rule based nature of the projections are stated directly."
    )
    doc.add_paragraph(
        "The most valuable next step is not additional visual complexity. It is validation with real revenue and operational records followed by honest out of sample testing. That would convert the current demonstration into a defensible business forecasting product."
    )

    add_table(
        doc,
        ["Overall assessment", "Finding"],
        [
            ("Portfolio value", "Strong demonstration of a complete data and analytics workflow"),
            ("Technical maturity", "Well structured prototype with tested core components"),
            ("Evidence maturity", "Requires validation on observed merchant and operational data"),
            ("Recommended positioning", "Decision support demonstration rather than production causal forecasting"),
        ],
        widths=[1.8, 4.9],
    )

    doc.add_heading("References and project files", level=1)
    refs = [
        "Project README and portfolio case study",
        "src data_generator.py data_pipeline.py econometrics.py revenue_analytics.py live_weather.py and revenue_connectors.py",
        "tests test_pipeline.py and test_connectors.py",
        "Open Meteo Weather Forecast API documentation",
        "NASA POWER Hourly API documentation",
        "NASA Worldview Earth observation browser",
    ]
    add_bullets(doc, refs)

    path = OUT / "weather_driven_business_intelligence_report.docx"
    doc.save(path)
    return path


def cue(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"Cue  {text}")
    r.bold = True
    r.font.color.rgb = RGBColor(23, 54, 93)
    return p


def spoken(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    r.font.size = Pt(11)
    return p


def build_script():
    doc = Document()
    configure_doc(doc, "Three Person Presentation Script")
    add_cover(
        doc,
        "Weather Driven Business Intelligence Presentation Script",
        "Three Presenter Version",
        "13 October 2026",
        "A coordinated 12 to 15 minute script with speaking lines handoffs demonstration cues and question preparation",
    )

    doc.add_heading("How to use this script", level=1)
    doc.add_paragraph(
        "Replace Presenter 1 Presenter 2 and Presenter 3 with your names. Rehearse once with the dashboard open and once without it. The script targets approximately 13 minutes and leaves time for questions. If the live weather service is unavailable use the fallback data and explain that the application is designed to remain demonstrable during an API outage."
    )
    add_table(
        doc,
        ["Presenter", "Primary responsibility", "Target time"],
        [
            ("Presenter 1", "Problem project overview and data engineering", "4 minutes"),
            ("Presenter 2", "Econometric method scenario logic and interpretation", "4 minutes"),
            ("Presenter 3", "Dashboard demonstration evaluation and conclusion", "4 to 5 minutes"),
        ],
        widths=[1.25, 4.25, 1.2],
    )

    doc.add_heading("Presenter 1 Opening and data foundation", level=1)
    cue(doc, "Slide 1  Title")
    spoken(
        doc,
        "Good morning. Our project is called the Weather Driven Business Intelligence Platform. It explores a practical business question: when weather changes customer behavior and operating conditions how can a company estimate the likely impact and respond before the disruption becomes expensive?",
    )
    spoken(
        doc,
        "Different businesses react in opposite ways to the same event. Heavy rain may increase food delivery and rideshare demand while reducing restaurant footfall bike usage and outdoor retail activity. Our platform brings those effects into one portfolio view instead of examining each business in isolation.",
    )
    cue(doc, "Slide 2  Project objective")
    spoken(
        doc,
        "The platform has three objectives. First it aligns weather and business records that arrive at different frequencies and in different timezones. Second it estimates weather relationships after accounting for normal calendar patterns. Third it translates the analysis into scenarios charts and operational guidance that a nontechnical manager can understand.",
    )
    spoken(
        doc,
        "An important point is that the historical dataset is synthetic. We designed it to demonstrate a realistic analytical workflow across delivery mobility retail hospitality energy and freight. We do not present the simulated results as verified effects from real companies.",
    )
    cue(doc, "Slide 3  Data sources and scale")
    spoken(
        doc,
        "The first raw source contains 69106 hourly weather observations. The second contains 279249 business records recorded every 15 minutes. The data cover New York City Chicago Seattle and Phoenix during 2023 and 2024.",
    )
    spoken(
        doc,
        "The sources deliberately do not line up. Weather timestamps are in UTC while business timestamps are local. Weather is hourly while business activity is every 15 minutes. Some weather readings are missing. Solving these mismatches is a central part of the project rather than an afterthought.",
    )
    cue(doc, "Slide 4  Data pipeline")
    spoken(
        doc,
        "The pipeline assigns each city its correct IANA timezone converts local business records to UTC aggregates business activity to hourly intervals and joins both sources by city and hour. It then creates calendar indicators rain categories weather lags temperature anomalies and cross sector ratios.",
    )
    spoken(
        doc,
        "The final analytical warehouse contains 70045 synchronized hourly observations and 61 variables. In our latest verification it contained no missing values and no duplicate city hour keys. This clean warehouse becomes the input to the modeling layer.",
    )
    cue(doc, "Handoff to Presenter 2")
    spoken(
        doc,
        "With the data foundation established Presenter 2 will explain how we separate weather relationships from recurring business patterns and how the scenario engine uses those estimates.",
    )

    doc.add_heading("Presenter 2 Method and business interpretation", level=1)
    cue(doc, "Slide 5  Why correlation can mislead")
    spoken(
        doc,
        "A raw correlation is not enough. Suppose delivery orders are high on a rainy Friday evening. Rain may contribute but Friday evenings are normally busy even in clear weather. If we ignore that calendar pattern we can attribute too much of the increase to rain.",
    )
    spoken(
        doc,
        "For each of 21 business metrics the project fits an ordinary least squares regression containing precipitation snow temperature anomaly and wind. It also controls for hour of day day of week month holiday status and city. Robust standard errors are used to reduce sensitivity to unequal residual variance.",
    )
    cue(doc, "Slide 6  Interpretation and boundaries")
    spoken(
        doc,
        "The result is a calendar and city adjusted estimate. For example a rain coefficient describes the modeled change associated with one additional millimeter of precipitation while the included controls remain fixed. We compare this adjusted estimate with the raw correlation to show how confounding can change an apparent relationship.",
    )
    spoken(
        doc,
        "We deliberately use the phrase adjusted association rather than proven causal effect. The data are simulated and the model does not create a natural experiment. R squared describes in sample model fit and a p value is evidence relative to a null hypothesis. Neither should be described as guaranteed predictive accuracy or certainty.",
    )
    cue(doc, "Slide 7  Cross sector scenario")
    spoken(
        doc,
        "The scenario engine lets a user change city weekday hour rain snow temperature anomaly and wind. It compares a clear weather baseline with the selected conditions across six sectors.",
    )
    spoken(
        doc,
        "Consider a heavy Friday evening storm. Bike usage may fall as commuters move toward transit and rideshare. More vehicles can increase congestion. Restaurant dine in demand may decline while delivery demand rises. At the same time courier availability may fall which can extend delivery times and increase cancellations. This is the domino effect the dashboard is designed to communicate.",
    )
    spoken(
        doc,
        "These chains are plausible scenario narratives rather than a separately estimated causal network. Their purpose is to help managers consider second order effects such as staffing packaging capacity courier incentives inventory positioning and freight delays.",
    )
    cue(doc, "Handoff to Presenter 3")
    spoken(
        doc,
        "Presenter 3 will now show how these ideas appear in the dashboard and summarize the current strengths limitations and next steps.",
    )

    doc.add_page_break()
    doc.add_heading("Presenter 3 Dashboard demonstration and conclusion", level=1)
    cue(doc, "Slide 8 or live dashboard  Portfolio controls")
    spoken(
        doc,
        "The application is built with Streamlit and Plotly. At the top we select one of four cities and choose a portfolio from more than 16 preset business categories. We can also create a custom business by specifying its baseline revenue and weather sensitivities.",
    )
    cue(doc, "Demonstration  Show Weekly Operational Impact")
    spoken(
        doc,
        "The first tab compares the modeled revenue series with a clear weather baseline and displays the net weather related shift. The historical tab expands the time horizon and includes the raw versus adjusted comparison discussed earlier.",
    )
    cue(doc, "Demonstration  Show Domino Chains")
    spoken(
        doc,
        "The interactions tab shows how a weather event may move demand between businesses and create downstream pressure. The descriptions change with the selected city businesses and weather regime.",
    )
    cue(doc, "Demonstration  Show Live Scenario Simulator")
    spoken(
        doc,
        "The live scenario tab uses Open Meteo for current and seven day weather conditions. Revenue projections then apply weekday patterns and configured business sensitivities. The guidance is rule based and transparent. It is not a black box AI prediction.",
    )
    spoken(
        doc,
        "NASA POWER is used as a recent reference source and NASA Worldview provides an external satellite imagery link. Because the NASA reference can be delayed it should not be interpreted as a same moment confirmation of current conditions.",
    )
    cue(doc, "Slide 9  Verification and limitations")
    spoken(
        doc,
        "The project includes seven automated checks covering data generation alignment scenario output connector schemas CSV detection and elasticity estimation. All seven tests passed in our latest run.",
    )
    spoken(
        doc,
        "The main limitation is that the default business history and the automatic Square and Stripe streams are simulated. The connector classes contain live API paths but the application does not provide credentials by default. A merchant can upload CSV or Excel history but those estimates still require enough overlapping dates and proper validation.",
    )
    cue(doc, "Slide 10  Conclusion")
    spoken(
        doc,
        "In conclusion this project demonstrates a complete analytical pipeline from mismatched raw feeds to business focused decision support. Its strengths are modular engineering explainable assumptions cross sector thinking and an accessible interface.",
    )
    spoken(
        doc,
        "The next step is to validate the approach with permissioned real merchant data report genuine uncertainty and backtest the forecasts. That would move the platform from a strong portfolio prototype toward a dependable business intelligence product. Thank you. We are happy to take questions.",
    )

    doc.add_heading("Fast version if time is reduced", level=1)
    doc.add_paragraph(
        "For an eight minute slot each presenter should keep the opening and closing paragraph of their section plus one technical paragraph. Skip the detailed NASA explanation and show only the portfolio selector scenario tab and diagnostics tab during the demonstration."
    )

    doc.add_heading("Likely questions and suggested answers", level=1)
    add_table(
        doc,
        ["Question", "Suggested answer"],
        [
            ("Is the data real", "The historical cross sector dataset is synthetic. Current weather can be live and users may upload real merchant revenue for exploratory calibration."),
            ("Why not use simple correlation", "Weather and commerce share calendar and seasonal patterns. The adjusted regression controls for several of those recurring patterns before estimating the weather relationship."),
            ("Does the model prove causality", "No. It estimates adjusted associations. Causal claims would require stronger identification and validation on observed data."),
            ("Where is AI used", "The present forecast and advice layers are deterministic and rule based. We avoid claiming a trained AI model where none is implemented."),
            ("Are Square and Stripe live", "The classes support credential based calls but the default application runs calibrated sandbox sequences. Production integration needs secure credentials pagination and monitoring."),
            ("How would you improve accuracy", "Use real merchant histories add holdout backtesting uncertainty intervals clustered or panel appropriate inference and forecast benchmarks."),
            ("What happens if an API fails", "The weather module returns a marked fallback package so the demonstration remains usable. The UI should always show whether data are live or fallback."),
        ],
        widths=[2.1, 4.6],
        font_size=8.8,
    )

    doc.add_heading("Rehearsal checklist", level=1)
    add_bullets(
        doc,
        [
            "Replace presenter labels with names and agree on pronunciation of econometrics and heteroskedasticity.",
            "Open the dashboard before presenting and confirm the city selector charts and scenario controls respond.",
            "Keep a screenshot or PDF of the main views available in case network access fails.",
            "Do not call the synthetic revenue real or the adjusted regression causal.",
            "Pause briefly at each handoff and let the next presenter begin before changing slides.",
            "Aim for one complete rehearsal tonight and one timed rehearsal tomorrow before presenting.",
        ],
    )

    path = OUT / "three_person_presentation_script.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    print(build_report())
    print(build_script())
